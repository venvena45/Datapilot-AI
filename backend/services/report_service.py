from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import io
import re
from datetime import datetime

def _clean_markdown_text(text: str) -> str:
    """Removes raw python list/dict dumps and cleans markdown artifacts."""
    if not text:
        return ""
    # Remove raw python dict/list reprs like (if [{'name': ...}])
    text = re.sub(r'\(if\s*\[\{.*?\}\]\)', '', text, flags=re.DOTALL)
    # Remove raw dict prints
    text = re.sub(r'\{[^{}]*:[^{}]*\}', '', text)
    return text.strip()

def _format_num(val) -> str:
    """Formats numeric values nicely with commas and 2 decimals if needed."""
    if val is None:
        return "-"
    try:
        num = float(val)
        if num == int(num) and abs(num) < 1e9:
            return f"{int(num):,}"
        return f"{num:,.2f}"
    except (ValueError, TypeError):
        return str(val)

def export_to_pdf(filename: str, rows_count: int, columns: list, stats: dict, insights: str, numeric_cols: list = None, cat_cols: list = None, dt_col: str = None) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36,
        title=f"Analytics Report - {filename}"
    )

    styles = getSampleStyleSheet()
    
    # Custom Brand Colors
    PRIMARY = colors.HexColor("#0F172A")    # Slate 900
    SECONDARY = colors.HexColor("#0284C7")  # Sky 600
    ACCENT = colors.HexColor("#10B981")     # Emerald 500
    TEXT_DARK = colors.HexColor("#1E293B")  # Slate 800
    TEXT_MUTED = colors.HexColor("#64748B") # Slate 500
    BG_LIGHT = colors.HexColor("#F8FAFC")   # Slate 50
    BORDER_CLR = colors.HexColor("#E2E8F0") # Slate 200

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=PRIMARY,
        alignment=TA_LEFT,
        spaceAfter=4
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=SECONDARY,
        spaceAfter=12
    )

    section_heading = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=PRIMARY,
        spaceBefore=14,
        spaceAfter=6,
        keepWithNext=True
    )

    sub_heading = ParagraphStyle(
        'SubHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=SECONDARY,
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=TEXT_DARK,
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        'BulletCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13.5,
        textColor=TEXT_DARK,
        leftIndent=14,
        firstLineIndent=-10,
        spaceAfter=4
    )

    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=TEXT_DARK
    )

    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white
    )

    story = []

    # 1. Title Banner
    story.append(Paragraph("AI Business Intelligence & Analytics Report", title_style))
    story.append(Paragraph("NexusAnalytics AI Data Studio", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=SECONDARY, spaceBefore=0, spaceAfter=10))

    # 2. Dataset Overview Summary Box
    meta_data = [
        [
            Paragraph(f"<b>Dataset File:</b> {filename or 'Unnamed Dataset'}", table_cell_style),
            Paragraph(f"<b>Generated:</b> {datetime.now().strftime('%Y-%m-%d %H:%M')}", table_cell_style),
        ],
        [
            Paragraph(f"<b>Total Records:</b> {rows_count:,}", table_cell_style),
            Paragraph(f"<b>Total Columns:</b> {len(columns)}", table_cell_style),
        ],
        [
            Paragraph(f"<b>Time Dimension:</b> {dt_col or 'N/A'}", table_cell_style),
            Paragraph(f"<b>Main Category:</b> {', '.join(cat_cols[:2]) if cat_cols else 'N/A'}", table_cell_style),
        ]
    ]
    meta_table = Table(meta_data, colWidths=[270, 270])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), BG_LIGHT),
        ('BOX', (0, 0), (-1, -1), 0.75, BORDER_CLR),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#F1F5F9")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 10))

    # 3. Statistical Metrics Table
    if stats and isinstance(stats, dict):
        story.append(Paragraph("Statistical Metrics Overview", section_heading))
        table_rows = [[
            Paragraph("Metric / Measure", table_header_style),
            Paragraph("Mean", table_header_style),
            Paragraph("Std Dev", table_header_style),
            Paragraph("Min", table_header_style),
            Paragraph("Max", table_header_style),
            Paragraph("Count", table_header_style),
        ]]

        valid_items = [
            (col, s) for col, s in stats.items()
            if not (col.lower() == 'id' or col.lower().endswith('_id') or col.lower().endswith(' id'))
        ]
        if not valid_items:
            valid_items = list(stats.items())

        for idx, (col_name, s) in enumerate(valid_items):
            if isinstance(s, dict):
                mean_str = _format_num(s.get("mean"))
                std_str = _format_num(s.get("std"))
                min_str = _format_num(s.get("min"))
                max_str = _format_num(s.get("max"))
                count_str = f"{s.get('count', rows_count):,}" if s.get('count') else f"{rows_count:,}"

                table_rows.append([
                    Paragraph(f"<b>{col_name}</b>", table_cell_style),
                    Paragraph(mean_str, table_cell_style),
                    Paragraph(std_str, table_cell_style),
                    Paragraph(min_str, table_cell_style),
                    Paragraph(max_str, table_cell_style),
                    Paragraph(count_str, table_cell_style),
                ])

        if len(table_rows) > 1:
            stat_table = Table(table_rows, colWidths=[150, 75, 75, 75, 85, 80])
            stat_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
                ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('GRID', (0, 0), (-1, -1), 0.5, BORDER_CLR),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('LEFTPADDING', (0, 0), (-1, -1), 6),
                ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BG_LIGHT]),
            ]))
            story.append(stat_table)
            story.append(Spacer(1, 12))

    # 4. Business Intelligence & Insights
    cleaned_insights = _clean_markdown_text(insights)
    if cleaned_insights:
        story.append(Paragraph("Business Intelligence & Key Findings", section_heading))
        lines = cleaned_insights.split("\n")
        
        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            # Format markdown bold **text** to ReportLab <b>text</b>
            formatted_line = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', line_str)
            # Format markdown italic *text* to ReportLab <i>text</i>
            formatted_line = re.sub(r'(?<!\*)\*(?!\*)(.*?)(?<!\*)\*(?!\*)', r'<i>\1</i>', formatted_line)

            if line_str.startswith("# "):
                # Top level header already exists in report
                continue
            elif line_str.startswith("## "):
                header_text = formatted_line.replace("## ", "").strip()
                story.append(Paragraph(header_text, sub_heading))
            elif line_str.startswith("### "):
                header_text = formatted_line.replace("### ", "").strip()
                story.append(Paragraph(header_text, sub_heading))
            elif line_str.startswith("- ") or line_str.startswith("* "):
                bullet_text = formatted_line[2:].strip()
                story.append(Paragraph(f"&bull; {bullet_text}", bullet_style))
            else:
                story.append(Paragraph(formatted_line, body_style))

    doc.build(story)
    return buffer.getvalue()


def export_to_docx(filename: str, rows_count: int, columns: list, stats: dict, insights: str, numeric_cols: list = None, cat_cols: list = None, dt_col: str = None) -> bytes:
    doc = Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Document Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("AI Business Intelligence & Analytics Report")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)
    title_p.paragraph_format.space_after = Pt(2)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("NexusAnalytics AI Data Studio")
    sub_run.font.size = Pt(10)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(2, 132, 199)
    sub_p.paragraph_format.space_after = Pt(14)

    # 2. Overview Metadata Box Table
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_info = [
        (f"Dataset File: {filename or 'Unnamed Dataset'}", f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"),
        (f"Total Records: {rows_count:,}", f"Total Columns: {len(columns)}"),
        (f"Time Dimension: {dt_col or 'N/A'}", f"Main Category: {', '.join(cat_cols[:2]) if cat_cols else 'N/A'}")
    ]

    for row_idx, (c1, c2) in enumerate(meta_info):
        row = meta_table.rows[row_idx]
        cell1, cell2 = row.cells[0], row.cells[1]
        cell1.text = c1
        cell2.text = c2
        cell1.paragraphs[0].runs[0].font.size = Pt(9)
        cell2.paragraphs[0].runs[0].font.size = Pt(9)
        
        # Light grey background
        shading1 = parse_xml(r'<w:shd {} w:fill="F8FAFC"/>'.format(nsdecls('w')))
        shading2 = parse_xml(r'<w:shd {} w:fill="F8FAFC"/>'.format(nsdecls('w')))
        cell1._tc.get_or_add_tcPr().append(shading1)
        cell2._tc.get_or_add_tcPr().append(shading2)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 3. Statistical Metrics Table
    if stats and isinstance(stats, dict):
        h1 = doc.add_heading(level=1)
        h1_run = h1.add_run("Statistical Metrics Overview")
        h1_run.font.size = Pt(14)
        h1_run.font.bold = True
        h1_run.font.color.rgb = RGBColor(15, 23, 42)

        valid_items = [
            (col, s) for col, s in stats.items()
            if not (col.lower() == 'id' or col.lower().endswith('_id') or col.lower().endswith(' id'))
        ]
        if not valid_items:
            valid_items = list(stats.items())

        stat_table = doc.add_table(rows=1, cols=6)
        stat_table.style = 'Table Grid'
        stat_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = ["Metric / Measure", "Mean", "Std Dev", "Min", "Max", "Count"]
        hdr_cells = stat_table.rows[0].cells
        for i, h_text in enumerate(headers):
            hdr_cells[i].text = h_text
            p = hdr_cells[i].paragraphs[0]
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(9)
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            shd = parse_xml(r'<w:shd {} w:fill="0F172A"/>'.format(nsdecls('w')))
            hdr_cells[i]._tc.get_or_add_tcPr().append(shd)

        for col_name, s in valid_items:
            if isinstance(s, dict):
                row_cells = stat_table.add_row().cells
                vals = [
                    col_name,
                    _format_num(s.get("mean")),
                    _format_num(s.get("std")),
                    _format_num(s.get("min")),
                    _format_num(s.get("max")),
                    f"{s.get('count', rows_count):,}" if s.get('count') else f"{rows_count:,}"
                ]
                for idx, v in enumerate(vals):
                    row_cells[idx].text = str(v)
                    p = row_cells[idx].paragraphs[0]
                    p.runs[0].font.size = Pt(9)
                    if idx == 0:
                        p.runs[0].font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 4. Insights & Strategic Recommendations
    cleaned_insights = _clean_markdown_text(insights)
    if cleaned_insights:
        h_insights = doc.add_heading(level=1)
        h_ins_run = h_insights.add_run("Business Intelligence & Key Findings")
        h_ins_run.font.size = Pt(14)
        h_ins_run.font.bold = True
        h_ins_run.font.color.rgb = RGBColor(15, 23, 42)

        lines = cleaned_insights.split("\n")
        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            if line_str.startswith("# "):
                continue
            elif line_str.startswith("## "):
                h2 = doc.add_heading(level=2)
                h2_text = line_str.replace("## ", "").replace("**", "").strip()
                h2_run = h2.add_run(h2_text)
                h2_run.font.size = Pt(11.5)
                h2_run.font.bold = True
                h2_run.font.color.rgb = RGBColor(2, 132, 199)
                h2.paragraph_format.space_before = Pt(8)
                h2.paragraph_format.space_after = Pt(3)
            elif line_str.startswith("- ") or line_str.startswith("* "):
                bullet_content = line_str[2:].strip()
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_after = Pt(3)
                # Parse bold parts
                parts = re.split(r'(\*\*.*?\*\*)', bullet_content)
                for part in parts:
                    if part.startswith("**") and part.endswith("**"):
                        run = p.add_run(part[2:-2])
                        run.font.bold = True
                    else:
                        p.add_run(part)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                parts = re.split(r'(\*\*.*?\*\*)', line_str)
                for part in parts:
                    if part.startswith("**") and part.endswith("**"):
                        run = p.add_run(part[2:-2])
                        run.font.bold = True
                    else:
                        p.add_run(part)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
